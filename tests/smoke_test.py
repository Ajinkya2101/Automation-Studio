"""End-to-end smoke test: upload -> record (simulated user) -> automated runs.

Run from the Prototype folder:  .venv\\Scripts\\python tests\\smoke_test.py
Uses port 8765 so it does not clash with a running studio.
"""
import asyncio
import io
import json
import os
import sys
import tempfile
import threading
import time
import urllib.request
import uuid
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
# Keep test suites and runs out of the studio's real data folder.
os.environ.setdefault("STUDIO_DATA_DIR", tempfile.mkdtemp(prefix="studio-smoke-"))

import openpyxl  # noqa: E402
import uvicorn  # noqa: E402

import demo_site  # noqa: E402
from app import storage  # noqa: E402
from app.events import sessions  # noqa: E402
from app.excel_io import build_workbook  # noqa: E402
from app.main import app  # noqa: E402
from app.recorder import record_session  # noqa: E402

PORT = 8765
BASE = f"http://127.0.0.1:{PORT}"
CASE_ID = "TS_SUP001"
NAME = "Northwind Traders"
NOTES = "Preferred supplier for office furniture. Payment terms 30 days."
app.include_router(demo_site.router)


def make_workbook() -> Path:
    """The test script a tester would upload, in the team template."""
    steps = [
        ("Sign in", "1. Open the sign-in page\n2. Input User Name & Password\n3. Click Sign in",
         {"URL": f"{BASE}/demo/login", "User Name": "demo.user", "Password": "demo123"},
         "User is signed in and the Suppliers list is displayed"),
        ("Open a new supplier", "1. Click Add supplier", {}, "New supplier form opens"),
        ("Enter the supplier", "1. Enter the Name\n2. Enter the Phone\n3. Enter the Notes",
         {"Name": NAME, "Phone": "+1 555 0100", "Notes": NOTES}, "Name, Phone and Notes are populated"),
        ("Save the supplier", "1. Click Save", {}, "Supplier saved confirmation is displayed"),
        ("Verify the supplier in the list", "1. Click Suppliers\n2. Confirm the supplier is listed",
         {"Name": NAME}, "Supplier appears in the list with the correct name"),
    ]
    case = {"id": CASE_ID, "sheet": CASE_ID, "title": "Create a supplier", "description": "", "steps": [
        {"step_id": f"{CASE_ID}.{i}", "name": name, "role": "User", "description": desc, "data": data,
         "expected": expected} for i, (name, desc, data, expected) in enumerate(steps, start=1)]}
    path = Path(tempfile.mkdtemp(prefix="studio-smoke-xlsx-")) / "Create_Supplier.xlsx"
    build_workbook(path, case, "Create a supplier")
    return path


def http(method, path, body=None, headers=None):
    req = urllib.request.Request(BASE + path, data=body, method=method, headers=headers or {})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read())


def upload(path: Path):
    boundary = uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{path.name}\"\r\n"
            f"Content-Type: application/octet-stream\r\n\r\n").encode() + path.read_bytes() + f"\r\n--{boundary}--\r\n".encode()
    return http("POST", "/api/suites", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})


def wait_idle(timeout=120):
    end = time.time() + timeout
    while time.time() < end:
        if not http("GET", "/api/status")["active"]:
            return
        time.sleep(0.5)
    raise TimeoutError("session did not finish")


def act_like_a_user(page):
    """What a tester would do in the recording window, including toolbar clicks."""
    def type_into(label, text):
        page.get_by_label(label, exact=True).click()
        page.keyboard.type(text, delay=5)

    def add_check(locator):
        page.locator("#check").click()
        locator.click()

    def next_step():
        page.locator("#next").click()
        page.wait_for_timeout(300)

    # Step 1: sign in
    type_into("User Name", "demo.user")
    type_into("Password", "demo123")
    page.get_by_role("button", name="Sign in").click()
    add_check(page.get_by_role("heading", name="Suppliers"))
    next_step()
    # Step 2: open the form
    page.get_by_role("link", name="Add supplier").click()
    add_check(page.get_by_role("heading", name="New supplier"))
    next_step()
    # Step 3: enter
    type_into("Name", NAME)
    type_into("Phone", "+1 555 0100")
    type_into("Notes", NOTES)
    next_step()
    # Step 4: save
    page.get_by_role("button", name="Save").click()
    add_check(page.get_by_text("Supplier saved", exact=True))
    next_step()
    # Step 5: verify in the list
    page.get_by_role("link", name="Suppliers", exact=True).click()
    add_check(page.locator(".name", has_text=NAME))
    page.locator("#next").click()  # "Finish recording"


def run_case(suite_id, case_id, data=None):
    demo_site.reset()  # every run starts from an empty list
    info = http("POST", f"/api/suites/{suite_id}/cases/{case_id}/run",
                json.dumps({"headed": False, "data": data or {}}).encode(),
                {"Content-Type": "application/json"})
    try:
        wait_idle()  # the failing run alone waits 45s for its check
    except TimeoutError:
        for ev in sessions.get(info["id"]).log.all():
            print("   log:", ev["type"], ev["message"][:400])
        raise
    if not storage.load_run(info["run_id"]):
        for ev in sessions.get(info["id"]).log.all():
            print("   log:", ev["type"], ev["message"][:400])
    return http("GET", f"/api/runs/{info['run_id']}")


def check_issue_log(run):
    """The failed run's issue log has what is needed to look into it on another computer."""
    failed = run["steps"][3]
    diag = failed.get("diagnostics") or {}
    check(run["environment"]["device"] and run["environment"]["browser"].get("user_agent"),
          f"run records its device and browser: {run['environment']['device']}")
    check(diag.get("page", {}).get("url", "").startswith(BASE) and diag.get("error", {}).get("traceback"),
          f"failed step keeps the page ({diag.get('page', {}).get('url')}) and the full traceback")
    check(run.get("trace") and (storage.run_dir(run["id"]) / "trace.zip").exists(), "Playwright trace kept for the failed run")

    for path, expect_runs in ((f"/api/runs/{run['id']}/issue-log", 1), ("/api/issue-log", 1)):
        with urllib.request.urlopen(BASE + path, timeout=60) as r:
            data = r.read()
            name = r.headers.get("content-disposition", "")
        z = zipfile.ZipFile(io.BytesIO(data))
        files = z.namelist()
        root = files[0].split("/")[0]
        has = lambda suffix: any(f.endswith(suffix) for f in files)
        check(all(has(s) for s in ("issues.xlsx", "README.txt", "environment.json", "app.log",
                                   f"runs/{run['id']}/log.txt", f"runs/{run['id']}/trace.zip",
                                   f"runs/{run['id']}/run.json", "recording.json", "case.json"))
              and any("/failure/" in f and f.endswith(".html") for f in files)
              and any("/sessions/" in f and "_record_" in f for f in files),
              f"{path} exports {len(files)} files as {name.split('filename=')[-1]}")
        ws = openpyxl.load_workbook(io.BytesIO(z.read(f"{root}/issues.xlsx")))["Issues"]
        kinds = [ws.cell(r, 4).value for r in range(2, ws.max_row + 1)]
        check("Step failed" in kinds, f"issues.xlsx lists the failure: {kinds}")
        text = z.read(f"{root}/runs/{run['id']}/log.txt").decode("utf-8")
        check("Full error from Playwright" in text and "Page:" in text, "log.txt has the page and full error")
        exported = z.read(f"{root}/runs/{run['id']}/run.json").decode("utf-8") + "".join(
            z.read(f).decode("utf-8") for f in files if f.endswith((".json", ".txt", ".log")) and "/runs/" not in f)
        check("demo123" not in exported, "the password is masked in everything exported as text")
        runs_ws = openpyxl.load_workbook(io.BytesIO(z.read(f"{root}/issues.xlsx")))["Runs"]
        check(runs_ws.max_row - 1 == expect_runs, f"{runs_ws.max_row - 1} run(s) in the export (passed runs left out)")


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        raise SystemExit(1)


def main():
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=PORT, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    while not server.started:
        time.sleep(0.1)

    up = upload(make_workbook())
    suite_id, case_id = up["suite_id"], up["cases"][0]
    check(case_id == CASE_ID and not up["warnings"], f"upload parsed {case_id} with no warnings")

    suite = storage.load_suite(suite_id)
    case = storage.find_case(suite, case_id)
    rec_session = sessions.start("record", record_session, suite_id=suite_id, case_id=case_id, suite=suite, case=case,
                   headless=True, on_ready=act_like_a_user)
    wait_idle()
    rec = storage.load_recording(suite_id, case_id)
    if not rec or rec["status"] != "completed":
        for ev in rec_session.log.all():
            print("   log:", ev["type"], ev["message"][:300])
    check(rec and rec["status"] == "completed", "recording completed")
    counts = [len(s["actions"]) for s in rec["steps"]]
    check(all(counts), f"every step has actions {counts}")
    params = sorted({a["param"] for s in rec["steps"] for a in s["actions"] if a.get("param")})
    check(params == ["Name", "Notes", "Password", "Phone", "User Name"], f"values linked to Excel data {params}")
    nav = [a["target"] for s in rec["steps"] for a in s["actions"] if a.get("navigates")]
    check(len(nav) >= 3, f"page-loading clicks detected: {nav}")

    run = run_case(suite_id, case_id)
    check(run["status"] == "passed", f"automated run passed ({run['passed']} steps, {run['duration_ms']}ms)")
    shots = list((storage.run_dir(run["id"]) / "screenshots").glob("*.png"))
    check(len(shots) == 5, f"{len(shots)} step screenshots saved")
    ws = openpyxl.load_workbook(storage.run_dir(run["id"]) / "results.xlsx")[CASE_ID]
    statuses = [ws.cell(r, 8).value for r in range(11, 16)]
    check(statuses == ["Pass"] * 5, f"results workbook statuses {statuses}")
    check(ws.cell(5, 2).value == "Automation (prototype)", "tester name written back")

    run2 = run_case(suite_id, case_id, {"Name": "Changed name from the dashboard"})
    check(run2["status"] == "passed", "same automation passes with a different Name")
    check("Changed name from the dashboard" in run2["steps"][4]["observed"],
          f"check followed the new data: {run2['steps'][4]['observed']}")

    run3 = run_case(suite_id, case_id, {"Phone": "not-a-number"})
    by = {s["step_id"]: s["status"] for s in run3["steps"]}
    check(run3["status"] == "failed" and by[f"{CASE_ID}.4"] == "failed" and by[f"{CASE_ID}.5"] == "skipped",
          f"bad data fails at the Save check and skips the rest: {by}")
    print("   failure logged as:", run3["steps"][3]["error"][:140])
    check_issue_log(run3)
    print("\nAll smoke checks passed.")
    server.should_exit = True


if __name__ == "__main__":
    main()
