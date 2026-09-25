"""Build the sample workbooks in the implementation team's template layout.

- Automation_Email.xlsx : TS_EML001 against the built-in Acme Mail demo app
- Outlook_Email.xlsx    : TS_OUT001 sends a real email with Outlook on the web
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).parent

ACME = dict(
    path=HERE / "Automation_Email.xlsx", case_id="TS_EML001", role="User",
    title="Automation email: create, write and send an email",
    description="Log in to Acme Mail, compose an email to the finance team, send it and "
                "confirm it appears in Sent",
    heading="Automation email (create, write and send)",
    steps=[
        ("Login to Acme Mail",
         "1. Open a browser and go to the Acme Mail login page\n2. Input User Name & Password\n3. Click Sign in",
         "User Name: demo.user\nPassword: demo123", "User is logged in and the Inbox is displayed"),
        ("Open a new message", "Navigation:\nInbox\nSteps:\n1. Click Compose", "", "New message form opens"),
        ("Write the email",
         "Navigation:\nContinue from previous step\nSteps:\n1. Enter the recipient in To\n"
         "2. Enter the Subject\n3. Enter the message body",
         "To: finance.team@acme.test\nSubject: Q3 inventory adjustment summary\n"
         "Body: Hi team, the Q3 inventory adjustments have been posted. Regards, Demo User",
         "To, Subject and Message are populated"),
        ("Send the email", "Navigation:\nContinue from previous step\nSteps:\n1. Click Send", "",
         "Message sent confirmation is displayed"),
        ("Verify the email in Sent",
         "Navigation:\nSent folder\nSteps:\n1. Click Sent\n2. Confirm the email is listed with the correct subject",
         "Subject: Q3 inventory adjustment summary", "Email appears in Sent with the correct subject"),
    ],
)

OUTLOOK = dict(
    path=HERE / "Outlook_Email.xlsx", case_id="TS_OUT001", role="Mailbox user",
    title="Outlook email: create, write and send a real email",
    description="Using Outlook on the web, send an email from abhosale@acsesolutions.com to "
                "abhosale@acsesolutions.com and confirm it appears in Sent Items. "
                "{now} in the Subject makes every run's email unique.",
    heading="Outlook email (create, write and send)",
    steps=[
        ("Open Outlook on the web",
         "Pre-requisite: use 'Sign in' in Automation Studio once (MFA included).\n"
         "Steps:\n1. Open Outlook on the web\n2. Wait for the mailbox to load",
         "URL: https://outlook.office.com/mail/", "The mailbox opens with the Inbox displayed"),
        ("Start a new email", "Navigation:\nInbox\nSteps:\n1. Click New mail", "",
         "A new, empty message opens"),
        ("Write the email",
         "Navigation:\nContinue from previous step\nSteps:\n1. Click To, enter the recipient and press Enter\n"
         "2. Click Subject and enter the subject\n3. Click in the message body and type the message",
         "To: abhosale@acsesolutions.com\nSubject: Automation Studio test {now}\n"
         "Body: This email was sent automatically by the Automation Studio prototype as part of test TS_OUT001.",
         "Recipient, Subject and message body are populated"),
        ("Send the email", "Navigation:\nContinue from previous step\nSteps:\n1. Click Send", "",
         "The message is sent and the new-message pane closes"),
        ("Verify the email in Sent Items",
         "Navigation:\nFolder pane\nSteps:\n1. Open Sent Items\n2. Confirm the email is listed with the subject",
         "Subject: Automation Studio test {now}", "The email appears in Sent Items with the correct subject"),
    ],
)


def build(path: Path, case_id: str, role: str, title: str, description: str, heading: str, steps) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = case_id
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    label_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="B7C3D0")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    first, last = 11, 10 + len(steps)
    ws["A1"], ws["F1"] = "Unit Test Case", "Summary"
    ws["G1"] = f'=IFERROR(COUNTIF($H${first}:$H${last},"Pass")/COUNTA($A${first}:$A${last}),0)'
    ws["G1"].number_format = "0%"
    ws["A1"].font = Font(bold=True, size=13)
    meta = [("Test Case ID", case_id), ("Test Case Title", title), ("Test Case Description", description),
            ("Tester Name", None), ("Test Location (Office)", None), ("Test Date & Time", None)]
    for i, (label, value) in enumerate(meta, start=2):
        ws.cell(i, 1, label).font = bold
        ws.cell(i, 1).fill = label_fill
        ws.cell(i, 2, value)
    ws["H6"], ws["H7"] = "Overall Coverage Status", "Overall Pass / Fail Status"

    headers = ["Test Step ID", "Process Step", "Role", "Test Step Description", "Test Data",
               "Expected Results", "Observed Results", "Status", "SR Number/Incident ID(s)",
               "Notes (e.g. Incident description)"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(9, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = box

    ws.cell(10, 1, case_id).font = bold
    ws.cell(10, 3, role)
    ws.cell(10, 4, heading).font = bold
    for i, (name, desc, data, expected) in enumerate(steps, start=1):
        r = 10 + i
        row = (f"{case_id}.{i}", name, role, desc, data or None, expected)
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.alignment = wrap
            cell.border = box
        for c in range(len(row) + 1, len(headers) + 1):
            ws.cell(r, c).border = box
            ws.cell(r, c).alignment = wrap

    for col, width in zip("ABCDEFGHIJ", (14, 26, 14, 50, 44, 34, 34, 10, 16, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A10"
    wb.save(path)
    print("written", path)


if __name__ == "__main__":
    for spec in (ACME, OUTLOOK):
        build(**spec)
