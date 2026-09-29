"""Read test scripts from the implementation team's Excel template and write
run results back into a copy of it.

Template (same layout as Test Script.xlsx): one sheet per test case, a header
row starting with "Test Step ID", the test-case row (e.g. TS_EML001) and then
the steps (TS_EML001.1, TS_EML001.2 ...). Test Data holds "Key: value" lines.
"""
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

COLUMN_ALIASES = {
    "step_id": ("test step id", "step id"),
    "name": ("fusion process step", "process step", "step name", "step"),
    "role": ("role",),
    "description": ("test step description", "description"),
    "test_data": ("test data",),
    "expected": ("expected results", "expected result"),
    "observed": ("observed results", "observed result"),
    "status": ("status",),
    "sr": ("sr number/incident id(s)",),
    "notes": ("notes (e.g. incident description)", "notes"),
}
META_LABELS = {
    "test case id": "id",
    "test case title": "title",
    "test case description": "description",
    "tester name": "tester",
    "test date & time": "tested_at",
}
KEY_LINE = re.compile(r"^\s*([A-Za-z][\w ./()&-]{0,40}?)\s*:\s*(.*)$")

STATUS_FILL = {
    "Pass": PatternFill("solid", fgColor="C6EFCE"),
    "Fail": PatternFill("solid", fgColor="FFC7CE"),
    "Not Run": PatternFill("solid", fgColor="EDEDED"),
}


class WorkbookError(ValueError):
    pass


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def _text(v) -> str:
    return str(v).strip() if v is not None else ""


def _broken(v) -> bool:
    """Formula remnants such as =#REF! are not usable values."""
    s = _text(v)
    return not s or s.startswith("=") or s.startswith("#")


def parse_test_data(raw: str) -> dict[str, str]:
    """'Key: value' per line; lines without a key continue the previous value."""
    data: dict[str, str] = {}
    last = None
    for line in _text(raw).splitlines():
        m = KEY_LINE.match(line)
        if m:
            last = m.group(1).strip()
            data[last] = m.group(2).strip()
        elif last and line.strip():
            data[last] = (data[last] + "\n" + line.strip()).strip()
    return data


def _find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60)):
        for cell in row:
            if _norm(cell.value) in COLUMN_ALIASES["step_id"]:
                cols = {}
                for c in ws[cell.row]:
                    n = _norm(c.value)
                    for key, aliases in COLUMN_ALIASES.items():
                        if n in aliases and key not in cols:
                            cols[key] = c.column
                return cell.row, cols
    return None, {}


def _meta_rows(ws, header_row: int) -> dict[str, int]:
    rows = {}
    for r in range(1, header_row):
        key = META_LABELS.get(_norm(ws.cell(r, 1).value))
        if key:
            rows[key] = r
    return rows


def parse_workbook(path: Path) -> tuple[list[dict], list[str]]:
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:
        raise WorkbookError(f"Could not open the workbook: {exc}") from exc

    cases, warnings = [], []
    for ws in wb.worksheets:
        header_row, cols = _find_header(ws)
        if not header_row:
            warnings.append(f"Sheet '{ws.title}': no 'Test Step ID' header row found, skipped.")
            continue
        missing = [k for k in ("step_id", "description") if k not in cols]
        if missing:
            warnings.append(f"Sheet '{ws.title}': missing column(s) {', '.join(missing)}, skipped.")
            continue

        def get(r, key):
            return _text(ws.cell(r, cols[key]).value) if key in cols else ""

        meta_rows = _meta_rows(ws, header_row)
        meta = {k: ws.cell(r, 2).value for k, r in meta_rows.items()}

        steps, case_heading = [], ""
        for r in range(header_row + 1, ws.max_row + 1):
            sid = get(r, "step_id")
            if not sid:
                continue
            if "." not in sid:  # the test-case row, e.g. "TS_INV004"
                case_heading = get(r, "description") or get(r, "name")
                continue
            raw_data = get(r, "test_data")
            data = parse_test_data(raw_data)
            steps.append({
                "step_id": sid,
                "row": r,
                "name": get(r, "name") or sid,
                "role": get(r, "role"),
                "description": get(r, "description"),
                "test_data_raw": raw_data,
                "data": data,
                "expected": get(r, "expected"),
            })
            empty = [k for k, v in data.items() if not v]
            if empty:
                warnings.append(f"{sid}: test data has no value for {', '.join(empty)}.")

        if not steps:
            warnings.append(f"Sheet '{ws.title}': no test steps found, skipped.")
            continue

        case_id = _text(meta.get("id"))
        if _broken(case_id):
            broken_value = _text(meta.get("id")) or "empty"
            case_id = steps[0]["step_id"].split(".")[0] or ws.title
            if meta_rows.get("id"):
                warnings.append(f"Sheet '{ws.title}': Test Case ID cell is broken "
                                f"({broken_value}); used {case_id} from the step IDs.")
        title = _text(meta.get("title"))
        if _broken(title):
            title = case_heading or ws.title
        cases.append({
            "id": case_id,
            "sheet": ws.title,
            "title": title,
            "description": "" if _broken(meta.get("description")) else _text(meta.get("description")),
            "steps": steps,
        })

    if not cases:
        raise WorkbookError("No test cases found. " + " ".join(warnings))
    return cases, warnings


TEMPLATE_HEADERS = ["Test Step ID", "Process Step", "Role", "Test Step Description", "Test Data",
                    "Expected Results", "Observed Results", "Status", "SR Number/Incident ID(s)",
                    "Notes (e.g. Incident description)"]


def data_text(data: dict[str, str]) -> str:
    return "\n".join(f"{k}: {v}" for k, v in data.items())


def build_workbook(path: Path, case: dict, heading: str = "") -> None:
    """Write a test case as a workbook in the team template.

    Used for the sample workbooks and for tests recorded without an Excel file,
    so every test has a script the team can read, edit and upload again.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = case["sheet"][:31]
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    label_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="B7C3D0")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    steps = case["steps"]
    role = steps[0].get("role", "") if steps else ""

    first, last = 11, 10 + max(len(steps), 1)
    ws["A1"], ws["F1"] = "Unit Test Case", "Summary"
    ws["G1"] = f'=IFERROR(COUNTIF($H${first}:$H${last},"Pass")/COUNTA($A${first}:$A${last}),0)'
    ws["G1"].number_format = "0%"
    ws["A1"].font = Font(bold=True, size=13)
    meta = [("Test Case ID", case["id"]), ("Test Case Title", case["title"]),
            ("Test Case Description", case.get("description") or None),
            ("Tester Name", None), ("Test Location (Office)", None), ("Test Date & Time", None)]
    for i, (label, value) in enumerate(meta, start=2):
        ws.cell(i, 1, label).font = bold
        ws.cell(i, 1).fill = label_fill
        ws.cell(i, 2, value)
    ws["H6"], ws["H7"] = "Overall Coverage Status", "Overall Pass / Fail Status"

    for c, h in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(9, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = box

    ws.cell(10, 1, case["id"]).font = bold
    ws.cell(10, 3, role or None)
    ws.cell(10, 4, heading or case["title"]).font = bold
    for r, s in enumerate(steps, start=first):
        row = (s["step_id"], s["name"], s.get("role") or None, s.get("description") or None,
               data_text(s.get("data", {})) or None, s.get("expected") or None)
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.alignment = wrap
            cell.border = box
        for c in range(len(row) + 1, len(TEMPLATE_HEADERS) + 1):
            ws.cell(r, c).border = box
            ws.cell(r, c).alignment = wrap

    for col, width in zip("ABCDEFGHIJ", (14, 26, 14, 50, 44, 34, 34, 10, 16, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A10"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_results(src: Path, dst: Path, sheet: str, run: dict) -> None:
    """Fill Observed Results / Status (and tester, date) in a copy of the workbook."""
    wb = openpyxl.load_workbook(src)
    ws = wb[sheet]
    header_row, cols = _find_header(ws)
    by_id = {s["step_id"]: s for s in run["steps"]}
    labels = {"passed": "Pass", "failed": "Fail"}

    for r in range(header_row + 1, ws.max_row + 1):
        sid = _text(ws.cell(r, cols["step_id"]).value)
        res = by_id.get(sid)
        if not res:
            continue
        status = labels.get(res["status"], "Not Run")
        if "observed" in cols:
            ws.cell(r, cols["observed"]).value = res.get("observed", "")
        if "status" in cols:
            cell = ws.cell(r, cols["status"])
            cell.value = status
            cell.fill = STATUS_FILL[status]
        if "notes" in cols and res.get("error"):
            ws.cell(r, cols["notes"]).value = res["error"]

    meta_rows = _meta_rows(ws, header_row)
    if "tester" in meta_rows:
        ws.cell(meta_rows["tester"], 2).value = "Automation (prototype)"
    if "tested_at" in meta_rows:
        ws.cell(meta_rows["tested_at"], 2).value = datetime.now().strftime("%d-%b-%Y %H:%M")
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
